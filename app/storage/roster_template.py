"""Excel 学生名单导入模板（生成器）。

模板固定三张工作表，**第一张就是导入时被读取的那张**：

============  ================================================================
``名单``      只有表头，教师从第 2 行开始填写。表头与
              :func:`app.storage.excel_io.auto_mapping` 的别名完全一致，
              导入时字段自动对应好，不需要手工做映射。
``填写说明``  每一列怎么填、可以自己加哪些列、常见问题。
``示例``      一份填好的样例数据，可整行复制到「名单」页再改成自己的数据。
============  ================================================================

教师也可以直接在程序里生成：菜单「文件 → 下载名单导入模板…」，
或左侧学生面板的「下载导入模板」按钮 —— 走的就是本模块。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Tuple

from .. import config
from .excel_io import OPENPYXL_AVAILABLE, ExcelError

if OPENPYXL_AVAILABLE:  # pragma: no branch - 环境相关
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- 常量
TEMPLATE_FILENAME = "学生名单导入模板.xlsx"

TEMPLATE_SHEET = "名单"
GUIDE_SHEET = "填写说明"
EXAMPLE_SHEET = "示例"

#: 模板列（也是建议的填写顺序）
COLUMNS: Tuple[str, ...] = ("学号", "姓名", "性别", "标签", "身高", "视力", "备注")

_COLUMN_WIDTHS: Tuple[int, ...] = (14, 11, 8, 20, 9, 9, 26)

#: 随包分发的模板文件（``scripts/make_templates.py`` 生成，可随时重新生成）
RESOURCE_TEMPLATE = config.TEMPLATES_DIR / TEMPLATE_FILENAME

#: 示例数据。学号故意以 ``X`` 开头，避免与真实学号冲突。
EXAMPLE_ROWS: Tuple[Tuple[Any, ...], ...] = (
    ("X1001", "张小明", "男", "班干部", 168, 4.8, "个子高，建议后排"),
    ("X1002", "李小红", "女", "优等生、班干部", 158, 5.1, ""),
    ("X1003", "王小刚", "男", "需关注", 172, 5.0, ""),
    ("X1004", "赵小美", "女", "视力差", 155, 4.2, "需要坐前排"),
    ("X1005", "陈小强", "男", "", 165, 4.9, ""),
    ("X1006", "周小雨", "女", "视力差、内向", 150, 4.3, "需要坐前排"),
)

EXAMPLE_NOTE = (
    "以上为示例数据（学号以 X 开头，不会和真实学号重复）。"
    "可以整行复制到「名单」工作表，再改成自己班级的数据；示例行本身不会被自动导入。"
)

#: 填写说明：(小标题, 正文)
GUIDE_SECTIONS: Tuple[Tuple[str, str], ...] = (
    (
        "三步用起来",
        "1. 在「名单」工作表里，从第 2 行开始逐行填写学生。\n"
        "2. 保存文件，回到程序：菜单「文件 → 导入学生名单…」，或点左侧面板的「导入 Excel」。\n"
        "3. 选中本文件即可。字段会自动对应好，直接确认即可。",
    ),
    (
        "表头不要改",
        "表头文字就是程序用来识别字段的标准名称，请不要修改表头，也不要删除整列。\n"
        "如果确实用了别的写法（如「学籍号」「身高(cm)」），导入时程序也能自动识别常见别名；\n"
        "识别不了的列，可以在导入对话框里手工指定它对应哪个字段。",
    ),
    (
        "学号（必填）",
        "唯一，不能重复。写成 20240101、1、A12 都可以，程序按文本处理。\n"
        "排序使用自然序：1-2-10 而不是 1-10-2。\n"
        "学号为空或与已有名单重复的行，导入时会被跳过并在提示里列出行号。",
    ),
    (
        "姓名（必填）",
        "姓名为空的行同样会被跳过。",
    ),
    (
        "性别（选填）",
        "填「男」或「女」，单元格带下拉可选。留空不影响导入。",
    ),
    (
        "标签（选填）",
        "一个学生可以有多个标签，用「、」或逗号分隔，例如：班干部、优等生。\n"
        "导入后程序会自动把这些标签建进标签库并分配颜色，可直接用于排位规则\n"
        "（如「视力差的学生必须坐前 3 排」）。",
    ),
    (
        "身高 / 视力（选填）",
        "填数字即可（身高按厘米）。它们是「数值属性」，可用于\n"
        "「按身高从矮到高、从前到后」排序、「成绩分 4 档均匀分布到各组」等规则。",
    ),
    (
        "备注（选填）",
        "任意文字，鼠标悬停在座位上时会显示。",
    ),
    (
        "可以自己加列",
        "除以上列外，还可以增加任意数字列（体重、成绩、年龄……）。\n"
        "新增列会自动成为一个数值属性，同样可以参与排位规则。",
    ),
    (
        "常见问题",
        "· 只想填学号和姓名？可以，其余列留空即可。\n"
        "· 60 人的班级？直接往下填 60 行，导入通常 1 秒内完成。\n"
        "· 模板丢了？在程序里随时重新生成：文件 → 下载名单导入模板…\n"
        "· 示例页的数据会被导入吗？不会，程序只读第一张「名单」表。",
    ),
)

_HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FF1F2430")
_HEADER_FILL = PatternFill("solid", fgColor="FFF2F4F7")
_THIN = Side(style="thin", color="FFD9DEE7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
_BODY_FONT = Font(name="微软雅黑", size=11)


def _require_openpyxl() -> None:
    if not OPENPYXL_AVAILABLE:
        raise ExcelError("未安装 openpyxl，无法生成 Excel 模板。请执行：pip install openpyxl")


def _write_header(sheet) -> None:
    """写入表头并设置列宽 / 冻结首行。"""
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _BORDER
    sheet.row_dimensions[1].height = 26
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_roster_sheet(sheet) -> None:
    """「名单」页：只有表头，留给教师填写。"""
    _write_header(sheet)
    sheet.freeze_panes = "A2"

    # 性别下拉。注意 OOXML 里 showDropDown=True 反而是“隐藏下拉”，故取 False；
    # showErrorMessage=False 让教师手写「男生」之类也不会被 Excel 拦住。
    validation = DataValidation(
        type="list",
        formula1='"男,女"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    sheet.add_data_validation(validation)
    validation.add("C2:C501")


def _write_guide_sheet(sheet) -> None:
    """「填写说明」页。"""
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 92

    sheet["A1"] = "学生名单导入模板 · 填写说明"
    sheet["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="FF2F6BFF")

    row = 3
    for title, body in GUIDE_SECTIONS:
        title_cell = sheet.cell(row=row, column=1, value=title)
        title_cell.font = _HEADER_FONT
        title_cell.alignment = _WRAP
        body_cell = sheet.cell(row=row, column=2, value=body)
        body_cell.font = _BODY_FONT
        body_cell.alignment = _WRAP
        sheet.row_dimensions[row].height = max(20, 16 * (body.count("\n") + 1) + 6)
        row += 1

    sheet.sheet_view.showGridLines = False


def _write_example_sheet(sheet) -> None:
    """「示例」页：一份填好的样例，可整行复制。"""
    _write_header(sheet)
    for values in EXAMPLE_ROWS:
        sheet.append(list(values))
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(COLUMNS)):
        for cell in row:
            cell.font = _BODY_FONT
            cell.alignment = _CENTER if cell.column != len(COLUMNS) else _LEFT
            cell.border = _BORDER

    note_row = len(EXAMPLE_ROWS) + 3
    sheet.cell(row=note_row, column=1, value=EXAMPLE_NOTE).font = Font(
        name="微软雅黑", size=10, color="FF6B7280"
    )
    sheet.freeze_panes = "A2"


def write_roster_template(path) -> Path:
    """生成名单导入模板；``path`` 没有 ``.xlsx`` 后缀时自动补上。"""
    _require_openpyxl()
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    if not target.parent.exists():
        target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    roster = workbook.active
    roster.title = TEMPLATE_SHEET
    _write_roster_sheet(roster)
    _write_guide_sheet(workbook.create_sheet(GUIDE_SHEET))
    _write_example_sheet(workbook.create_sheet(EXAMPLE_SHEET))
    try:
        workbook.save(str(target))
    except Exception as exc:  # noqa: BLE001 - 面向用户的可读错误
        raise ExcelError("生成模板失败：%s" % exc) from exc
    finally:
        workbook.close()
    return target
